from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_a4_margins(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

def configure_styles(doc):
    # Normal style
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    
    # Heading 1 (Chapter Title)
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = None
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.page_break_before = True
    
    # Heading 2 (Chapter Heading)
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.color.rgb = None
    h2.paragraph_format.line_spacing = 1.5
    
    # Table Caption
    if 'Caption' in doc.styles:
        caption = doc.styles['Caption']
        caption.font.name = 'Times New Roman'
        caption.font.size = Pt(12)
        caption.font.italic = False
        caption.font.color.rgb = None

def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def configure_headers_footers(doc):
    section = doc.sections[0]
    
    # Header
    header = section.header
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("SAMVED HealthGrid: Federated Disease Surveillance")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(10)
    
    # Footer
    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    # Use tabs to separate left and right
    fp.text = "Dept. of IT, RIT, Rajaramnagar\t\t\t\t\t\t"
    fp.runs[0].font.name = 'Times New Roman'
    fp.runs[0].font.size = Pt(10)
    
    run2 = fp.add_run()
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(10)
    add_page_number(run2)

def build_front_pages(doc):
    # COVER PAGE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n\n\nFINAL YEAR PROJECT REPORT\nON\n\n").bold = True
    
    title = p.add_run("SAMVED HealthGrid: Privacy-Preserving Federated Disease Surveillance and Early Outbreak Detection System\n\n")
    title.bold = True
    title.font.size = Pt(16)
    
    p.add_run("SUBMITTED BY:\n")
    p.add_run("1. Abhijeet Ravindra Gaikwad, PRN 2310053\n")
    p.add_run("2. Shubham Sanjay Shewalkar, PRN 2310033\n")
    p.add_run("3. Aditya Jagtap, PRN 2306123\n")
    p.add_run("4. Abhishek Waghmare, PRN 2310067\n\n")
    
    p.add_run("UNDER THE GUIDANCE OF:\nProf. P. T. Sawant\n\n")
    
    p.add_run("[Insert institute logo]\n\n")
    p.add_run("Department of Information Technology\n")
    p.add_run("Rajarambapu Institute of Technology, Rajaramnagar\n")
    p.add_run("Academic Year: [2025-2026]")
    
    # CERTIFICATE
    doc.add_page_break()
    doc.add_paragraph("CERTIFICATE", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("This is to certify that the project entitled \"SAMVED HealthGrid: Privacy-Preserving Federated Disease Surveillance and Early Outbreak Detection System\" has been successfully completed by Abhijeet Ravindra Gaikwad (PRN 2310053), Shubham Sanjay Shewalkar (PRN 2310033), Aditya Jagtap (PRN 2306123), and Abhishek Waghmare (PRN 2310067) under the guidance of Prof. P. T. Sawant in partial fulfillment of the requirements for the Degree of Bachelor of Technology in Information Technology.")
    doc.add_paragraph("\n\n[Signature]\nProf. P. T. Sawant\n(Project Guide)\n")
    doc.add_paragraph("\n[Signature]\n[HOD Name]\n(Head, Department of Information Technology)\n")
    doc.add_paragraph("\n[Signature]\n[Director Name]\n(Director)")
    
    # DECLARATION
    doc.add_page_break()
    doc.add_paragraph("DECLARATION", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("We, the undersigned, declare that the project entitled \"SAMVED HealthGrid: Privacy-Preserving Federated Disease Surveillance and Early Outbreak Detection System\" is a record of our original work carried out under the supervision of Prof. P. T. Sawant. The system implemented is a simulation-based proof of concept and has not been submitted elsewhere for the award of any degree or diploma.")
    doc.add_paragraph("\n1. Abhijeet Ravindra Gaikwad, PRN 2310053\n2. Shubham Sanjay Shewalkar, PRN 2310033\n3. Aditya Jagtap, PRN 2306123\n4. Abhishek Waghmare, PRN 2310067\n\nDate: [Insert Date]")

    # HACKATHON
    doc.add_page_break()
    doc.add_paragraph("SMC HACKATHON 2025–26 PARTICIPATION", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("The concept for this project originated from the \"Smart Public Health Management System\" problem statement presented by the Solapur Municipal Corporation during the SMC Hackathon 2025–26. Our team participated in the prototype round and subsequently extended the concept into this final-year academic project.\n\n[Insert Hackathon Participation Certificate Placeholder]")
    
    # ACKNOWLEDGEMENT
    doc.add_page_break()
    doc.add_paragraph("ACKNOWLEDGEMENT", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("We express our profound gratitude to our project guide, Prof. P. T. Sawant, for providing technical direction and continuous encouragement throughout this endeavor. We also thank the Solapur Municipal Corporation for inspiring the initial problem statement. Finally, we acknowledge the support of the Department of Information Technology at Rajarambapu Institute of Technology.")

    # ABSTRACT
    doc.add_page_break()
    doc.add_paragraph("ABSTRACT", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("Traditional public-health surveillance systems rely primarily on delayed, fragmented case reporting from isolated hospitals and laboratories, resulting in significant reporting delays that inhibit proactive outbreak management. This project addresses the limitation by proposing SAMVED HealthGrid, a simulation-based proof of concept that uses data-local geographic Federated Learning for multi-source syndromic surveillance and early outbreak-risk prediction. The system integrates simulated real-time data ingestion from hospitals, laboratories, pharmacies, and weather conditions (temperature, humidity, rainfall) across five geographic zones. Rather than transferring sensitive raw patient records to a centralized server, raw data remains at the zone client. A PyTorch Multi-Layer Perceptron (MLP) is trained locally at each zone, and the model parameters are aggregated centrally using the Federated Averaging (FedAvg) algorithm. The evaluation relies on a stochastic network SEIRS generator capable of simulating gradual epidemic waves and realistic observational delays. The system successfully predicts outbreak onset within a seven-day horizon and visualizes zone-level risk using a Geographic Information System (GIS) dashboard. While the system demonstrates the mathematical feasibility of decentralized surveillance without direct data sharing, a major observed limitation is the absolute burden of false warnings inherently tied to high-sensitivity early detection thresholds for rare outbreak events.")

    # TOC, LOF, LOT, LOA
    doc.add_page_break()
    doc.add_paragraph("TABLE OF CONTENTS", style='Heading 1').paragraph_format.page_break_before = False
    add_toc(doc)
    
    doc.add_page_break()
    doc.add_paragraph("LIST OF FIGURES", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("[Right-click and Update Field to generate List of Figures]")
    
    doc.add_page_break()
    doc.add_paragraph("LIST OF TABLES", style='Heading 1').paragraph_format.page_break_before = False
    doc.add_paragraph("[Right-click and Update Field to generate List of Tables]")
    
    doc.add_page_break()
    doc.add_paragraph("LIST OF ABBREVIATIONS AND ACRONYMS", style='Heading 1').paragraph_format.page_break_before = False
    abbrev = "API: Application Programming Interface\nFedAvg: Federated Averaging\nFL: Federated Learning\nGIS: Geographic Information System\nMLP: Multi-Layer Perceptron\nPR-AUC: Precision-Recall Area Under Curve\nSEIRS: Susceptible-Exposed-Infectious-Recovered-Susceptible"
    doc.add_paragraph(abbrev)

def build_chapter1(doc):
    doc.add_paragraph("CHAPTER 1: INTRODUCTION", style='Heading 1')
    doc.add_paragraph("1.1 Background and Context", style='Heading 2')
    doc.add_paragraph("Disease surveillance is a critical function of public-health authorities. Conventional systems rely heavily on manual reporting of confirmed cases from hospitals and laboratories. This reactive approach introduces substantial reporting delays, meaning outbreaks are frequently detected only after they have caused significant community transmission. Recently, there has been immense interest in syndromic surveillance, which uses early indicators such as pharmacy sales and weather conditions to detect anomalies before clinical confirmation. Concurrently, privacy constraints restrict the sharing of raw healthcare data across institutional boundaries, necessitating privacy-preserving architectures like Federated Learning [1].")
    
    doc.add_paragraph("1.2 Motivation", style='Heading 2')
    doc.add_paragraph("Early syndromic indicators from hospitals, laboratories, pharmacies, and weather stations contain complementary signals that can identify an outbreak trajectory days earlier than case counts alone. However, aggregating these heterogeneous data sources into a centralized database introduces significant privacy risks and administrative barriers. A decentralized approach allows regions to collaborate on training a global predictive model while maintaining data sovereignty.")
    
    doc.add_paragraph("1.3 Problem Statement", style='Heading 2')
    doc.add_paragraph("To design, implement, and evaluate a simulation-based proof of concept for a federated early disease outbreak detection system that aggregates multi-source syndromic and meteorological signals across geographic zones without centralizing raw institutional data.")
    
    doc.add_paragraph("1.4 Objectives", style='Heading 2')
    doc.add_paragraph("1. To design a multi-source data ingestion pipeline incorporating simulated hospital, laboratory, pharmacy, and weather data.\n2. To develop and evaluate a PyTorch Multi-Layer Perceptron for predicting outbreak onset within a seven-day future horizon.\n3. To implement data-local geographic Federated Learning using the FedAvg algorithm across five distinct zones.\n4. To evaluate model performance comparing local, federated, and centralized paradigms using a synthetic SEIRS-based dataset.\n5. To visualize zone-level outbreak risk using an interactive GIS dashboard.")
    
    doc.add_paragraph("1.5 Scope", style='Heading 2')
    doc.add_paragraph("The implemented system is a simulation-based proof of concept. The scope includes hospital symptom and admission signals, laboratory testing metrics, pharmacy medicine-demand signals, and weather indicators (temperature, humidity, rainfall). The architecture features five geographic zones acting as Federated Learning clients, PyTorch MLP local models, Flower FedAvg orchestration, SEIRS-based synthetic epidemic time-series generation, and a GIS-based zone-level risk visualization dashboard. The scope strictly excludes ICU-bed forecasting, oxygen-bed forecasting, ventilator forecasting, individual patient diagnosis, telemedicine, clinical deployment, differential privacy, and secure aggregation.")
    
    doc.add_paragraph("1.6 Users and Stakeholders", style='Heading 2')
    doc.add_paragraph("The intended stakeholders include public-health authorities, municipal decision-makers, researchers, and system administrators who require regional risk summaries rather than individual patient records.")
    
    doc.add_paragraph("1.7 Project Contributions", style='Heading 2')
    doc.add_paragraph("The project contributes a complete, reproducible simulation pipeline for evaluating multi-source syndromic surveillance under federated constraints. It provides an end-to-end integration of a SEIRS data generator, an automated evaluation pipeline, a Node.js backend, a FastAPI prediction service, and a React-based spatial dashboard.")
    
    doc.add_paragraph("1.8 Organization of the Report", style='Heading 2')
    doc.add_paragraph("Chapter 1 introduces the background and objectives. Chapter 2 analyzes the existing process and problem definition. Chapter 3 surveys the literature. Chapter 4 defines system requirements. Chapter 5 details the architectural design and database schemas. Chapter 6 explains the implementation of federated models and data generation. Chapter 7 covers testing. Chapter 8 provides deployment instructions. Chapter 9 presents the evaluation metrics. Chapter 10 concludes the report and discusses future scope.")
    doc.add_paragraph("Summary: This chapter outlined the context, motivation, scope, and objectives of developing a federated disease surveillance system.")

def build_chapter2(doc):
    doc.add_paragraph("CHAPTER 2: PROBLEM IDENTIFICATION AND ANALYSIS", style='Heading 1')
    doc.add_paragraph("2.1 Problem Identification", style='Heading 2')
    doc.add_paragraph("The problem was identified through the \"Smart Public Health Management System\" challenge presented by the Solapur Municipal Corporation during the SMC Hackathon 2025–26. The initial problem highlighted the administrative difficulties municipal bodies face when attempting to detect localized disease spread across disconnected wards and hospitals.")
    
    doc.add_paragraph("2.2 Problem Selection", style='Heading 2')
    doc.add_paragraph("The problem was selected because it intersects software engineering with societal health challenges. Developing a federated early-warning system provides extensive learning opportunities in machine learning, distributed systems, spatial databases, and full-stack development, making it an ideal final-year project.")
    
    doc.add_paragraph("2.3 Problem Definition", style='Heading 2')
    doc.add_paragraph("The system must generate daily and weekly surveillance summaries and predict the probability of an outbreak onset within the next seven days at a geographic zone level. The constraints are that raw institutional data must remain within the local zone, the evaluation must be fully simulated due to the lack of linked public healthcare records, and the system must provide a user-friendly GIS map for authorities.")
    
    doc.add_paragraph("2.4 Problem Analysis", style='Heading 2')
    doc.add_paragraph("A Root Cause Analysis identified that reporting delay is fundamentally caused by the requirement for definitive laboratory confirmation before an event is counted. Stakeholder Analysis revealed that hospitals are reluctant to share raw data due to privacy policies. A Fishbone Diagram categorization shows that technological causes involve disconnected electronic health record silos, while data causes involve the lack of real-time syndromic standardization.")
    
    doc.add_paragraph("2.5 Existing Process or Current System", style='Heading 2')
    doc.add_paragraph("In the conventional process, a patient visits a hospital, a sample is sent to a laboratory, and days later a confirmed positive result is reported to municipal authorities. Authorities aggregate these confirmed reports weekly, visually inspect the counts, and manually declare an outbreak when the incidence is surging.")
    
    doc.add_paragraph("2.6 Limitations of the Existing System", style='Heading 2')
    doc.add_paragraph("The existing system is inherently reactive. Because it waits for clinical confirmation, the reporting delay often exceeds the incubation period of the disease. It ignores early behavioral indicators like sudden surges in over-the-counter pharmacy sales for fever medication. Furthermore, geographic visibility is limited, and aggregating sensitive data centrally poses severe privacy risks.")
    doc.add_paragraph("Summary: This chapter analyzed the root causes of surveillance delays and justified the need for a federated, multi-source syndromic approach.")

def main():
    doc = Document()
    set_a4_margins(doc)
    configure_styles(doc)
    configure_headers_footers(doc)
    
    build_front_pages(doc)
    build_chapter1(doc)
    build_chapter2(doc)
    
    doc.save('temp_report_1.docx')

if __name__ == '__main__':
    main()
