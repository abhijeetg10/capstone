from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def create_report():
    doc = Document()
    
    # 1. Page size: A4
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    
    # Margins: Left 0.75, Right 0.5, Top 0.5, Bottom 0.5 inch
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(0)
    style_normal.paragraph_format.space_before = Pt(0)

    # Custom styles
    # Chapter Title: 14pt, bold, centered
    chapter_style = doc.styles.add_style('ChapterTitle', WD_STYLE_TYPE.PARAGRAPH)
    chapter_style.font.name = 'Times New Roman'
    chapter_style.font.size = Pt(14)
    chapter_style.font.bold = True
    chapter_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter_style.paragraph_format.line_spacing = 1.5
    
    # Heading inside chapter: 12pt, bold
    h2_style = doc.styles.add_style('Heading2', WD_STYLE_TYPE.PARAGRAPH)
    h2_style.font.name = 'Times New Roman'
    h2_style.font.size = Pt(12)
    h2_style.font.bold = True
    h2_style.paragraph_format.line_spacing = 1.5

    def add_chapter(title):
        doc.add_page_break()
        doc.add_paragraph(title, style='ChapterTitle')
        doc.add_paragraph()

    def add_h2(text):
        doc.add_paragraph(text, style='Heading2')
        
    def add_para(text):
        doc.add_paragraph(text, style='Normal')
        
    # Content Generation
    add_para("SAMVED HealthGrid: Privacy-Preserving Federated Disease Surveillance and Early Outbreak Detection System")
    add_para("Rajarambapu Institute of Technology, Rajaramnagar")
    add_para("Department of Information Technology")
    add_para("\nProject team:")
    add_para("1. Abhijeet Ravindra Gaikwad, PRN 2310053")
    add_para("2. Shubham Sanjay Shewalkar, PRN 2310033")
    add_para("3. Aditya Jagtap, PRN 2306123")
    add_para("4. Abhishek Waghmare, PRN 2310067")
    add_para("\nProject guide:\nProf. P. T. Sawant")
    
    # We will write the full MD instead to convert via pandoc or markdown2 for ease.
    # Because writing 30 pages of text through python-docx calls line by line is prone to error and truncation.
    doc.save('SAMVED_HealthGrid_Final_Project_Report.docx')

if __name__ == "__main__":
    create_report()
