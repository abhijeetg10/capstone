from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def append_diagrams():
    doc = Document('SAMVED_HealthGrid_Final_Project_Report.docx')
    
    # We will append them to the end of the document before references, 
    # but the easiest is just adding an Appendix.
    doc.add_page_break()
    doc.add_paragraph("APPENDIX A: ADDITIONAL DIAGRAMS", style='Heading 1')
    
    diagrams_to_add = [
        ("diagrams_output/class.png", "Figure A.1: Class Diagram"),
        ("diagrams_output/sequence.png", "Figure A.2: Sequence Diagram"),
        ("diagrams_output/activity.png", "Figure A.3: Activity Diagram"),
        ("diagrams_output/component.png", "Figure A.4: Component Diagram"),
        ("diagrams_output/dfd_level0.png", "Figure A.5: DFD Level 0"),
        ("diagrams_output/er.png", "Figure A.6: Entity Relationship (ER) Diagram")
    ]
    
    for path, caption in diagrams_to_add:
        if os.path.exists(path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=Inches(5.0))
            doc.add_paragraph(caption, style='Caption').alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            
    doc.save('SAMVED_HealthGrid_Final_Project_Report_v2.docx')

if __name__ == '__main__':
    append_diagrams()
